"""Verify the two generated tables."""
import collections
import glob
import os

from openpyxl import load_workbook

for path in ["output/SUPPLEMENT_sessions_tutorials.xlsx", "output/NO_DATA_courses.xlsx"]:
    wb = load_workbook(path)
    print("=" * 92)
    print(os.path.basename(path))
    for n in wb.sheetnames:
        ws = wb[n]
        print(f"   {n:26s} rows={ws.max_row:4d} cols={ws.max_column}")
    for n in wb.sheetnames:
        if n == "说明":
            continue
        ws = wb[n]
        print(f"\n   --- {n} header + 4 rows ---")
        for r in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            print("     " + " | ".join("" if v is None else str(v)[:26] for v in r))
    print()

# ---- cross-check the supplement against raw json ----
import json
supp = {}
for p in glob.glob("data_supplement/*.json"):
    b = json.load(open(p, encoding="utf-8"))
    for c in b["courses"]:
        supp[f"{b['subject']} {c['catalog_nbr']}"] = c

print("=" * 92)
print("spot checks")
for code in ["BIO 1001", "ECO 2011", "MAT 1005", "FIN 3080"]:
    c = supp.get(code)
    if not c:
        print(f"   {code}: not in supplement")
        continue
    ss = c.get("sections") or []
    print(f"   {code}: career={c.get('career')} sections={len(ss)} "
          f"terms={sorted({s['term'] for s in ss})}")
    for s in ss[:3]:
        print(f"        {s['session_code']:12s} {s['days']:6s} {s['start']}-{s['end']} "
              f"| {s['room'][:26]:26s} | {s['instructor'][:24]}")

# ---- no-data breakdown ----
wb = load_workbook("output/NO_DATA_courses.xlsx")
ws = wb["无数据课程"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
print(f"\nno-data rows: {len(rows)}")
print("by reason :", dict(collections.Counter(r[5] for r in rows)))
print("by college:", dict(collections.Counter(r[2] for r in rows).most_common(12)))
