"""Export scraped SIS data to Excel - one workbook per term.

    python export_excel.py --subject CSC
    python export_excel.py --data-dir data --outdir output
"""
import argparse
import glob
import json
import os
import re
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BAND = PatternFill("solid", fgColor="F2F7FB")

SESSION_COLS = [
    ("课程代码", 12), ("课程名称", 46), ("Session序号", 10), ("Session编号", 13),
    ("Class Nbr", 10), ("周几", 9), ("时间段", 13), ("地点 (Room)", 26),
    ("老师 (Instructor)", 30), ("类型", 8), ("开课状态", 10),
    ("上课日期", 24), ("Career", 14),
]
TUTORIAL_COLS = [
    ("课程代码", 12), ("课程名称", 46), ("Session序号", 10), ("Session编号", 13),
    ("Class Nbr", 10), ("周几", 9), ("时间段", 13), ("地点 (Room)", 26),
    ("老师 (Instructor)", 30), ("开课状态", 10), ("上课日期", 24),
]
ENROLL_COLS = [
    ("课程代码", 12), ("课程名称", 46), ("Career", 14), ("学分 (Units)", 11),
    ("Prerequisite 先修", 44), ("Exclusion 互斥", 40),
    ("Enrollment Requirement (原始)", 60),
]
COURSE_COLS = [
    ("课程代码", 12), ("课程名称", 46), ("Career", 14), ("学分 (Units)", 11),
    ("开课学院 (Academic Group)", 30), ("授课语言 (Attributes)", 26),
    ("课程简介 (Description)", 90),
]


def comp_type(session_code):
    m = re.search(r"-([A-Z]+)$", session_code or "")
    return m.group(1) if m else ""


def fmt_time(s):
    """'08:30'+'09:50' -> '08:30-09:50'; handles TBA / missing halves."""
    a = (s.get("start") or "").strip()
    b = (s.get("end") or "").strip()
    if a and b:
        return f"{a}-{b}"
    return a or b or ""


def is_tutorial(session_code):
    """TUT components are tutorials; everything else (LEC/LAB/SEM/...) is a session."""
    t = comp_type(session_code)
    if t:
        return t.startswith("TUT")
    return (session_code or "").upper().startswith("T")


def safe_title(t):
    """Excel sheet names cannot contain : \\ / ? * [ ] and max 31 chars."""
    return re.sub(r"[:\\/?*\[\]]", "-", t)[:31]


def write_sheet(ws, cols, rows, band=True):
    for c, (name, width) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 30
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=cols[c - 1][1] >= 40,
                                       horizontal="center" if cols[c - 1][1] <= 13 else "left")
            if band and r % 2 == 0:
                cell.fill = BAND
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"


def load_courses(data_dir, subjects=None):
    courses = []
    for path in sorted(glob.glob(os.path.join(data_dir, "*.json"))):
        code = os.path.splitext(os.path.basename(path))[0]
        if subjects and code not in subjects:
            continue
        with open(path, encoding="utf-8-sig") as f:
            blob = json.load(f)
        courses.extend(blob.get("courses", []))
    return courses


def build(term_code, term_label, courses, outdir, prefix):
    ug = [c for c in courses if c.get("career") == "Undergraduate" and not c.get("error")]

    sess_rows, tut_rows, enroll_rows, course_rows = [], [], [], []
    for c in sorted(ug, key=lambda x: (x.get("subject", ""), str(x.get("catalog_nbr", "")))):
        name = c.get("header") or f"{c.get('course')} - {c.get('list_title', '')}"
        subj = c.get("subject", "")
        enroll = c.get("enrollment") or {}
        enroll_rows.append([c.get("course"), name, c.get("career"), c.get("units"),
                            enroll.get("prerequisite", ""), enroll.get("exclusion", ""),
                            enroll.get("raw", "")])
        course_rows.append([c.get("course"), name, c.get("career"), c.get("units"),
                            c.get("academic_group", ""), c.get("attributes", ""),
                            c.get("description", "")])
        for s in c.get("sections") or []:
            if s.get("term") != term_label:
                continue
            base = [c.get("course"), name, s.get("session_no"), s.get("session_code"),
                    s.get("class_nbr"), s.get("days"), fmt_time(s),
                    s.get("room")]
            if is_tutorial(s.get("session_code")):
                tut_rows.append(base + [s.get("instructor"), s.get("status"), s.get("dates")])
            else:
                sess_rows.append(base + [s.get("instructor"), comp_type(s.get("session_code")),
                                         s.get("status"), s.get("dates"), c.get("career")])

    wb = Workbook()
    wb.remove(wb.active)
    for title, cols, rows in [("Sessions (课程安排)", SESSION_COLS, sess_rows),
                              ("Tutorials (辅导课)", TUTORIAL_COLS, tut_rows),
                              ("Enrollment (先修与互斥)", ENROLL_COLS, enroll_rows),
                              ("Courses (全部UG课程)", COURSE_COLS, course_rows)]:
        ws = wb.create_sheet(safe_title(title))
        write_sheet(ws, cols, rows)

    ws = wb.create_sheet(safe_title("说明 README"))
    info = [
        ("数据来源", "CUHK-Shenzhen SIS - Browse Course Catalog (ESTABLISH_COURSES.SSS_BROWSE_CATLG)"),
        ("学期 Term", term_label),
        ("Term Code", term_code),
        ("课程范围", f"{prefix} 前缀，Career = Undergraduate"),
        ("UG 课程数", len(ug)),
        ("Session 行数 (含 LEC/LAB 等)", len(sess_rows)),
        ("Tutorial 行数 (TUT)", len(tut_rows)),
        ("导出时间", __import__("time").strftime("%Y-%m-%d %H:%M:%S")),
        ("", ""),
        ("Sessions", "非 Tutorial 的教学安排：课程代码-课程名称-Session序号-周几-时间段-地点-老师"),
        ("Tutorials", "辅导课 TUT 安排：课程代码-课程名称-Session序号-周几-时间段-地点（另附老师）"),
        ("Enrollment", "先修 Prerequisite 与互斥 Exclusion，取自 Enrollment Requirement"),
        ("Courses", "全部本科课程的基本信息与课程简介"),
        ("", ""),
        ("说明", "Session序号 为 PeopleSoft 的 session number；Session编号 形如 L01-LEC / T01-TUT，用于区分同一课程的不同班次。"),
        ("说明", "时间段已统一为 24 小时制 HH:MM-HH:MM。"),
        ("说明", "仅包含该课程在所选学期确实开课的班次；未开课的学期不在本表中。"),
    ]
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    ws["A1"] = f"{prefix} 课程数据导出说明"
    ws["A1"].font = TITLE_FONT
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")

    os.makedirs(outdir, exist_ok=True)
    safe = term_label.replace(" ", "_")
    path = os.path.join(outdir, f"{prefix}_SIS_{safe}.xlsx")
    wb.save(path)
    return path, len(sess_rows), len(tut_rows), len(enroll_rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default=os.path.join(HERE, "data"))
    ap.add_argument("--outdir", default=os.path.join(HERE, "output"))
    ap.add_argument("--subject", default="CSC", help="comma separated subject codes")
    ap.add_argument("--all", action="store_true",
                    help="export every subject present in --data-dir")
    ap.add_argument("--prefix", default=None, help="output filename prefix")
    args = ap.parse_args()

    if args.all:
        subjects = sorted(os.path.splitext(os.path.basename(p))[0]
                          for p in glob.glob(os.path.join(args.data_dir, "*.json")))
    else:
        subjects = [s.strip().upper() for s in args.subject.split(",") if s.strip()]
    prefix = args.prefix or ("ALL" if (args.all or len(subjects) != 1) else subjects[0])
    courses = load_courses(args.data_dir,
                           None if args.all else set(subjects))
    print(f"loaded {len(courses)} courses from {len(subjects)} subjects")

    terms = OrderedDict([("2510", "2025-26 Term 1"), ("2520", "2025-26 Term 2"),
                         ("2610", "2026-27 Term 1")])
    for code, label in terms.items():
        path, ns, nt, ne = build(code, label, courses, args.outdir, prefix)
        print(f"  {label:18s} -> {os.path.basename(path)}  "
              f"sessions={ns} tutorials={nt} enrollment={ne}")


if __name__ == "__main__":
    main()
