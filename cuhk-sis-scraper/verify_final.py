"""Final verification of the delivered workbooks."""
import collections
import glob
import os
import re

from openpyxl import load_workbook

for path in sorted(glob.glob("output/*.xlsx")):
    wb = load_workbook(path)
    print("=" * 92)
    print(os.path.basename(path))
    for n in wb.sheetnames:
        ws = wb[n]
        print(f"   {n:28s} rows={ws.max_row:4d} cols={ws.max_column}")

    sess = wb["Sessions (课程安排)"]
    tut = wb["Tutorials (辅导课)"]
    enr = wb["Enrollment (先修与互斥)"]

    def col(ws, idx):
        return [r[idx] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

    subs = collections.Counter(str(c).split()[0] for c in col(sess, 0))
    tsubs = collections.Counter(str(c).split()[0] for c in col(tut, 0))
    print(f"   Sessions by subject : {dict(sorted(subs.items()))}")
    print(f"   Tutorials by subject: {dict(sorted(tsubs.items()))}")

    ai = sorted({c for c in col(enr, 0) if str(c).startswith("AIE")})
    print(f"   AIE courses in Enrollment: {ai}")
    ece = sorted({c for c in col(enr, 0) if str(c).startswith("ECE")})
    print(f"   ECE courses: {ece}")

    # component-type sanity
    types = collections.Counter(r[9] for r in sess.iter_rows(min_row=2, values_only=True) if r[0])
    print(f"   Session component types: {dict(types)}")
    ttypes = collections.Counter(re.sub(r"^\w+?-", "", str(r[3])) for r in
                                 tut.iter_rows(min_row=2, values_only=True) if r[0])
    print(f"   Tutorial component types: {dict(ttypes)}")
    print()
