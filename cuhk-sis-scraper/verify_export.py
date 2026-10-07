"""Read back the exported workbooks and verify structure + sample content."""
import glob
import os
import sys

from openpyxl import load_workbook

outdir = sys.argv[1] if len(sys.argv) > 1 else "output"
for path in sorted(glob.glob(os.path.join(outdir, "*.xlsx"))):
    wb = load_workbook(path)
    print("=" * 90)
    print(os.path.basename(path))
    print("  sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"    {name:26s} dims={ws.dimensions:14s} rows={ws.max_row:4d} cols={ws.max_column}")
    for sheet in ("Sessions (课程安排)", "Tutorials (辅导课)", "Enrollment (先修与互斥)"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        print(f"\n  --- {sheet} (first 4 data rows) ---")
        for r in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), values_only=True):
            print("    " + " | ".join("" if v is None else str(v)[:30] for v in r))
    print()
